"""Excel 解析 + 逐行比对 + 状态分类 + 导出。"""
import os
import uuid
from datetime import date
from openpyxl import load_workbook, Workbook

from flask import current_app
from ..extensions import db
from ..models import (
    Manufacturer, Product, BusinessLicense, FoodLicense, ThirdPartyReport, FactoryReport,
)
from ..utils.date_utils import status_of, status_label


def _threshold_days():
    from ..models import Setting
    s = Setting.query.first()
    return (s.default_threshold_days if s else 30) or 30


def parse_excel(file_storage):
    """解析 Excel 文件，返回 [[厂家名, 商品名, 规格], ...]。

    约定列：第一行表头（厂家名/商品名/规格），其后为数据行。
    """
    wb = load_workbook(file_storage, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    iterator = ws.iter_rows(values_only=True)
    first = next(iterator, None)
    if first is None:
        return rows
    # 若首行包含汉字表头则跳过
    if not _looks_like_header(first):
        rows.append(_normalize_row(first))
    for r in iterator:
        row = _normalize_row(r)
        if any(row):
            rows.append(row)
    wb.close()
    return rows


def _looks_like_header(row):
    head = " ".join(str(c) for c in row if c)
    return any(k in head for k in ["厂家", "商品", "规格", "名称"])


def _normalize_row(row):
    name = str(row[0]).strip() if len(row) > 0 and row[0] else ""
    product = str(row[1]).strip() if len(row) > 1 and row[1] else ""
    spec = str(row[2]).strip() if len(row) > 2 and row[2] else ""
    return [name, product, spec]


def compare_rows(rows):
    """对每行进行库内比对，返回结果列表。

    每行结果：{manufacturer_name, product_name, spec, found_manufacturer, found_product,
              bl_status, fl_status, tp_status, fc_status, tp_remaining, fc_remaining}
    """
    threshold = _threshold_days()
    today = date.today()
    results = []
    for name, product_name, spec in rows:
        result = {
            "manufacturer_name": name,
            "product_name": product_name,
            "spec": spec,
            "found_manufacturer": False,
            "found_product": False,
            "bl_status": "missing",
            "fl_status": "missing",
            "tp_status": "missing",
            "fc_status": "missing",
            "tp_remaining": None,
            "fc_remaining": None,
        }
        m = Manufacturer.query.filter_by(name=name).first()
        if not m:
            results.append(result)
            continue
        result["found_manufacturer"] = True
        # 厂家级证件
        if m.business_license:
            result["bl_status"] = status_of(m.business_license.valid_until, threshold, today)
        if m.food_license:
            result["fl_status"] = status_of(m.food_license.valid_until, threshold, today)
        # 商品级证件
        p = Product.query.filter_by(manufacturer_id=m.id, name=product_name, spec=spec or None).first()
        if not p:
            # 规格可能为空字符串匹配 None，再尝试一次
            p = Product.query.filter(
                Product.manufacturer_id == m.id,
                Product.name == product_name,
                db.or_(Product.spec == (spec or None), Product.spec.is_(None)),
            ).first()
        if not p:
            results.append(result)
            continue
        result["found_product"] = True
        # 最新第三方报告
        latest_tp = (
            ThirdPartyReport.query.filter_by(product_id=p.id)
            .order_by(ThirdPartyReport.test_date.desc())
            .first()
        )
        if latest_tp:
            result["tp_status"] = status_of(latest_tp.valid_until, threshold, today)
            if latest_tp.valid_until:
                result["tp_remaining"] = (latest_tp.valid_until - today).days
        # 最新厂检（已取消到期日概念，仅记录是否存在）
        latest_fc = (
            FactoryReport.query.filter_by(product_id=p.id)
            .order_by(FactoryReport.test_date.desc())
            .first()
        )
        if latest_fc:
            result["fc_remaining"] = None
            result["fc_status"] = "normal"
        results.append(result)
    return results


def export_results(results, task_id):
    """将比对结果导出为 Excel，写入 uploads/exports/，返回相对路径。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "比对结果"
    ws.append([
        "厂家名", "商品名", "规格",
        "厂家存在", "商品存在",
        "营业执照", "食品经营许可证",
        "第三方检验报告", "第三方剩余天数",
        "出厂检验报告", "厂检剩余天数",
    ])
    for r in results:
        ws.append([
            r["manufacturer_name"], r["product_name"], r["spec"],
            "是" if r["found_manufacturer"] else "否",
            "是" if r["found_product"] else "否",
            status_label(r["bl_status"]),
            status_label(r["fl_status"]),
            status_label(r["tp_status"]),
            r["tp_remaining"] if r["tp_remaining"] is not None else "",
            status_label(r["fc_status"]),
            r["fc_remaining"] if r["fc_remaining"] is not None else "",
        ])
    # 调整列宽
    for col_idx, col_name in enumerate(["厂家名", "商品名", "规格"], start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = 18
    rel_dir = current_app.config["EXPORTS_SUBDIR"]
    abs_dir = os.path.join(current_app.config["UPLOAD_FOLDER"], rel_dir)
    os.makedirs(abs_dir, exist_ok=True)
    filename = f"compare_{task_id}.xlsx"
    abs_path = os.path.join(abs_dir, filename)
    wb.save(abs_path)
    wb.close()
    return os.path.join(rel_dir, filename)
